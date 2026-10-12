"""
Import the Excel void tracker (the Voids and TEMP tabs).

Columns are found by their heading, not their position. Cell colours are read
using the colour key in the top rows of each tab:

    Works Completed     -> trade complete
    Works in Progress   -> trade in progress
    Snagging / Follow On (yellow) -> snagging
    Issue (red fill)    -> issue
    Glazing / Comp Door (light blue) -> booked / ordered
    Hand Back This Week (lavender, on the address) -> no import value; worked out from dates
    Electrical Flag (purple) -> electrical flag
    Red writing         -> claimed for payment

Rows are matched to existing voids by SIMPRO number, or by address when there
is no SIMPRO number, so importing the same workbook again updates rather than
duplicates.
"""

import io
import re
from datetime import date, datetime, timedelta

import openpyxl
from django.db import transaction
from django.utils import timezone

from audit.models import record

from . import services
from .models import Area, JobType, Supervisor, Trade, TradeTask, Void, VoidImportRun, normalise_heading

S = TradeTask.Status
ISSUE = "issue"  # an Issue colour on a trade raises an issue on the whole job

FIELD_HEADINGS = {
    "postcode": {"post code", "postcode"},
    "supervisor": {"supervisor"},
    "area": {"void area", "area"},
    "simpro": {"simpro number", "simpro no", "simpro", "temp simpro number"},
    "expected": {"expected completion", "hand back date", "handback date", "hand back"},
    "address": {"address"},
    "access": {"access"},
    "status_note": {"status"},
    "last_inspection": {"last inspection date", "last inspection"},
}

LEGEND = {
    "works completed": S.COMPLETE,
    "works in progress": S.IN_PROGRESS,
    "snagging issued": S.SNAGGING,
    "follow on": S.SNAGGING,
    "issue": ISSUE,
    "glazing comp door": S.BOOKED,
    "electrical flag": "electrical",
    "hand back this week": "handback",
}
# Used when a tab has no colour key of its own (standard Excel colours).
DEFAULT_COLOURS = {
    "rgb:FF00B050": S.COMPLETE, "rgb:FF92D050": S.COMPLETE, "rgb:FFFFC000": S.IN_PROGRESS,
    "rgb:FFFFFF00": S.SNAGGING, "rgb:FFFF0000": ISSUE, "rgb:FF7030A0": "electrical",
}
RED_TEXT = {"FFFF0000", "FFC00000"}
DATE_IN_TEXT = re.compile(r"(\d{1,2})\s*[/.-]\s*(\d{1,2})(?:\s*[/.-]\s*(\d{2,4}))?")


class TrackerError(Exception):
    pass


def colour_key(cell):
    f = cell.fill
    if f is None or f.fill_type in (None, "none"):
        return ""
    c = f.fgColor
    if c is None:
        return ""
    if c.type == "rgb" and isinstance(c.rgb, str):
        return "" if c.rgb in ("00000000", "FFFFFFFF") else f"rgb:{c.rgb}"
    if c.type == "theme":
        return f"theme:{c.theme}:{round(c.tint or 0, 2)}"
    return ""


def is_red_text(cell):
    c = cell.font.color if cell.font else None
    return bool(c is not None and c.type == "rgb" and isinstance(c.rgb, str) and c.rgb in RED_TEXT)


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"[ \t\xa0]+", " ", str(value)).strip()


def guess_date(text, today):
    """'25/09' -> the closest sensible date (no year: within the last ~10 months or next 2)."""
    m = DATE_IN_TEXT.search(text or "")
    if not m:
        return None
    d, mo, y = int(m.group(1)), int(m.group(2)), m.group(3)
    try:
        if y:
            y = int(y)
            return date(y + 2000 if y < 100 else y, mo, d)
        for year in (today.year, today.year - 1, today.year + 1):
            cand = date(year, mo, d)
            if today - timedelta(days=300) <= cand <= today + timedelta(days=65):
                return cand
    except ValueError:
        return None
    return None


def to_date(value, today):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return guess_date(clean(value), today) if value else None


WORK_COLOURS = (S.COMPLETE, S.IN_PROGRESS, S.SNAGGING, ISSUE, S.BOOKED)


def parse_task(cell, legend, today):
    text = clean(cell.value)
    colour = legend.get(colour_key(cell)) or DEFAULT_COLOURS.get(colour_key(cell))
    claimed = bool(text) and is_red_text(cell)
    if not text and colour not in WORK_COLOURS:
        return None
    low = text.lower()
    task = {"status": S.STC, "claimed": claimed, "assignee": "", "date": None, "notes": ""}
    if low.startswith("stc"):
        task["status"] = S.STC
        rest = text[3:].strip(" -–:")
        task["notes"] = rest
        if colour in WORK_COLOURS:
            task["status"] = colour
        return task
    if low.startswith(("not needed", "not required", "n/a")):
        task["status"] = S.NOT_REQUIRED
        return task
    task["status"] = colour if colour in WORK_COLOURS else S.IN_PROGRESS
    lines = [ln.strip() for ln in re.split(r"[\n\r]+", text) if ln.strip()]
    last = lines[-1] if lines else ""
    d = guess_date(last, today)
    task["date"] = d.isoformat() if d else None
    m = DATE_IN_TEXT.search(last)
    who = (last[:m.start()] if m else last).strip(" -–:,")
    simple = len(lines) == 1 and m and not last[m.end():].strip(" -–") and len(who) <= 40
    if simple or (not m and len(lines) == 1 and len(who) <= 25):
        task["assignee"] = who
    else:
        task["assignee"] = who[:40] if who and len(who) <= 40 else ""
        task["notes"] = text
    return task


def read_legend(ws, header_row):
    legend = {}
    for row in ws.iter_rows(min_row=1, max_row=max(1, header_row - 1)):
        for cell in row:
            meaning = LEGEND.get(normalise_heading(cell.value))
            key = colour_key(cell)
            if meaning and key:
                legend[key] = meaning
    return legend


def find_header_row(ws):
    for row in ws.iter_rows(min_row=1, max_row=15):
        names = {normalise_heading(c.value) for c in row}
        if "address" in names and any("simpro" in n for n in names):
            return row[0].row
    raise TrackerError(f"Tab '{ws.title}': could not find the heading row (with Address and SIMPRO).")


def parse_workbook(data, today=None):
    """Read the workbook. Returns (rows, report)."""
    today = today or timezone.localdate()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # noqa: BLE001 - any unreadable file
        raise TrackerError("This file could not be read as an Excel workbook (.xlsx).") from exc
    job_types = {jt.sheet_name.strip().lower(): jt.name for jt in JobType.objects.exclude(sheet_name="")}
    trades = list(Trade.objects.filter(active=True))
    rows, report = [], {"sheets": [], "problems": [], "skipped_tabs": []}
    for ws in wb.worksheets:
        job_type = job_types.get(ws.title.strip().lower())
        if not job_type:
            report["skipped_tabs"].append(ws.title)
            continue
        header_row = find_header_row(ws)
        legend = read_legend(ws, header_row)
        fields, trade_cols, unknown = {}, {}, []
        for cell in ws[header_row]:
            h = normalise_heading(cell.value)
            if not h:
                continue
            field = next((f for f, names in FIELD_HEADINGS.items() if h in names), None)
            trade = next((t for t in trades if h in t.heading_names()), None)
            if field and field not in fields:
                fields[field] = cell.column - 1
            elif trade:
                trade_cols[trade.name] = cell.column - 1
            else:
                unknown.append(clean(cell.value))
        if "address" not in fields:
            raise TrackerError(f"Tab '{ws.title}': no Address column.")
        if unknown:
            report["problems"].append(f"{ws.title}: columns not recognised and ignored: {', '.join(unknown)}")
        count = 0
        for row in ws.iter_rows(min_row=header_row + 1):
            def val(f):
                return row[fields[f]] if f in fields and fields[f] < len(row) else None
            addr_cell = val("address")
            address = clean(addr_cell.value) if addr_cell is not None else ""
            simpro = clean(val("simpro").value) if val("simpro") is not None else ""
            filled = [c for c in row if clean(c.value)]
            if not address:
                if filled and simpro:
                    report["problems"].append(f"{ws.title} row {row[0].row}: SIMPRO {simpro} has no address – skipped.")
                continue
            if not simpro and len(filled) == 1:
                continue  # a divider row such as "NEW SYSTEM"
            item = {"sheet": ws.title, "row": row[0].row, "job_type": job_type, "simpro": simpro,
                    "address": address, "tasks": {}, "issue_note": ""}
            for f in ("postcode", "supervisor", "access", "status_note"):
                c = val(f)
                item[f] = clean(c.value) if c is not None else ""
            c = val("area")
            item["area"] = clean(c.value).upper() if c is not None else ""
            c = val("expected")
            item["expected"], item["expected_note"] = None, ""
            if c is not None and c.value is not None:
                if isinstance(c.value, (datetime, date)):
                    item["expected"] = to_date(c.value, today).isoformat()
                else:
                    item["expected_note"] = clean(c.value)
            c = val("last_inspection")
            d = to_date(c.value, today) if c is not None else None
            item["last_inspection"] = d.isoformat() if d else None
            row_colours = {legend.get(colour_key(x)) or DEFAULT_COLOURS.get(colour_key(x))
                           for x in row[: max(fields.values()) + 1]}
            item["electrical_flag"] = "electrical" in row_colours
            item["issue_flag"] = colour_key(addr_cell) in {k for k, v in legend.items() if v == ISSUE} \
                or colour_key(addr_cell) == "rgb:FFFF0000"
            for name, idx in trade_cols.items():
                if idx < len(row):
                    t = parse_task(row[idx], legend, today)
                    if t:
                        item["tasks"][name] = t
            issue_trades = [n for n, t in item["tasks"].items() if t["status"] == ISSUE]
            for n in issue_trades:
                item["tasks"][n]["status"] = S.IN_PROGRESS
            if issue_trades:
                item["issue_flag"] = True
                item["issue_note"] = "Issue on " + ", ".join(issue_trades) + " (from tracker)"
            if item["electrical_flag"] and "Electrical" not in item["tasks"]:
                # Flagged purple in the tracker: put Electrical on the job (the flag is automatic now).
                item["tasks"]["Electrical"] = {"status": S.STC, "claimed": False, "assignee": "", "date": None,
                                               "notes": "Electrical flag in tracker"}
            if "Electrical" in item["tasks"] and item["tasks"]["Electrical"]["status"] not in ("", S.NOT_REQUIRED):
                if re.search(r"rewire|electric", " ".join(str(v) for v in item["tasks"]["Electrical"].values()), re.I):
                    item["electrical_flag"] = True
            rows.append(item)
            count += 1
        report["sheets"].append({"name": ws.title, "job_type": job_type, "rows": count})
    if not report["sheets"]:
        raise TrackerError("None of the tabs in this workbook are set up for import (see Voids settings → Job types).")
    return merge_duplicates(rows, report), report


def row_key(item):
    if item["simpro"]:
        return ("simpro", item["simpro"].replace(" ", "").lower())
    return ("address", re.sub(r"[^a-z0-9]", "", item["address"].lower()))


def merge_duplicates(rows, report):
    """The same void on two rows (e.g. a scope note row): fill gaps from the later row."""
    out, seen = [], {}
    for item in rows:
        key = (item["job_type"], row_key(item))
        if key in seen:
            first = seen[key]
            for k, v in item.items():
                if k == "tasks":
                    for t, task in v.items():
                        first["tasks"].setdefault(t, task)
                elif k == "status_note" and v and v != first[k]:
                    first[k] = "; ".join(x for x in (first[k], v) if x)
                elif v and not first.get(k):
                    first[k] = v
            report["problems"].append(
                f"{item['sheet']} row {item['row']}: same void as row {first['row']} – merged into one.")
            continue
        seen[key] = item
        out.append(item)
    return out


def find_existing(item):
    qs = Void.objects.filter(job_type__name=item["job_type"])
    kind, value = row_key(item)
    if kind == "simpro":
        for v in qs.exclude(simpro=""):
            if v.simpro.replace(" ", "").lower() == value:
                return v
        return None
    for v in qs.filter(simpro=""):
        if re.sub(r"[^a-z0-9]", "", v.address.lower()) == value:
            return v
    return None


def find_supervisor(name, cache):
    n = (name or "").strip().lower()
    if not n:
        return None
    for s in cache:
        if n in s.names():
            return s
    first = n.split()[0]
    matches = [s for s in cache if s.name.lower().split()[0] == first]
    return matches[0] if len(matches) == 1 else None


def preview_upload(filename, data, user, today=None):
    rows, report = parse_workbook(data, today)
    sups = list(Supervisor.objects.all())
    unknown_sups = sorted({r["supervisor"] for r in rows if r["supervisor"] and not find_supervisor(r["supervisor"], sups)})
    if unknown_sups:
        report["problems"].append("Supervisor names not recognised (add them, or add the spelling to a supervisor's "
                                  "import names, then upload again): " + ", ".join(unknown_sups))
    no_sup = [f"{r['sheet']} row {r['row']} ({r['address']})" for r in rows
              if not r["postcode"] and not find_supervisor(r["supervisor"], sups)]
    if no_sup:
        report["problems"].append(f"{len(no_sup)} voids have no postcode or supervisor, so no area or supervisor "
                                  "can be set (add them on the site after importing): " + "; ".join(no_sup))
    report["new"] = sum(1 for r in rows if not find_existing(r))
    report["updated"] = len(rows) - report["new"]
    report["total"] = len(rows)
    return VoidImportRun.objects.create(uploaded_by=user, filename=filename[:255], rows=rows, report=report)


TRACKED = ["simpro", "address", "postcode", "area_id", "supervisor_id", "access", "handback_date",
           "status_note", "electrical_flag", "issue_flag", "last_inspection"]


@transaction.atomic
def apply_run(run, user):
    if run.status != VoidImportRun.Status.PREVIEW:
        raise TrackerError("This import has already been applied or discarded.")
    sups = list(Supervisor.objects.all())
    areas = {a.code: a for a in Area.objects.all()}
    job_types = {j.name: j for j in JobType.objects.all()}
    trades = {t.name: t for t in Trade.objects.all()}
    created = updated = 0
    for item in run.rows:
        void = find_existing(item)
        is_new = void is None
        if is_new:
            void = Void(job_type=job_types[item["job_type"]], imported=True)
        before = Void.objects.get(pk=void.pk) if not is_new else None
        void.simpro = item["simpro"]
        void.address = item["address"][:200]
        void.postcode = item["postcode"][:10].upper()
        void.access = item["access"][:100]
        # The tracker's "Expected Completion" is the client hand-back date.
        void.handback_date = date.fromisoformat(item["expected"]) if item["expected"] else None
        note = "; ".join(x for x in (item["status_note"], item["expected_note"]) if x)
        void.status_note = note[:300]
        void.electrical_flag = item["electrical_flag"]
        void.issue_flag = item["issue_flag"]
        if item["issue_flag"] and not void.issue_note:
            void.issue_note = item.get("issue_note", "")[:300]
        if item["last_inspection"]:
            void.last_inspection = date.fromisoformat(item["last_inspection"])
        if item["area"] and item["area"] in areas:
            void.area = areas[item["area"]]
        sup = find_supervisor(item["supervisor"], sups)
        if sup:
            void.supervisor = sup
        services.apply_postcode(void)
        if is_new:
            void.status = Void.Status.STC
        void.save()
        for name, t in item["tasks"].items():
            TradeTask.objects.update_or_create(void=void, trade=trades[name], defaults={
                "status": t["status"], "claimed": t["claimed"], "assignee": t["assignee"][:100],
                "date": date.fromisoformat(t["date"]) if t["date"] else None, "notes": t["notes"]})
        void.tasks.exclude(trade__name__in=list(item["tasks"])).delete()
        services.refresh_status(void, user, activity=False)
        if is_new:
            created += 1
            record(user, "void.import", void, f"Imported {void} from {item['sheet']} row {item['row']}")
        else:
            changes = services.changed_fields(before, void, TRACKED)
            updated += 1
            record(user, "void.import", void, f"Updated {void} from {item['sheet']} row {item['row']}", changes)
    run.status = VoidImportRun.Status.APPLIED
    run.applied_at = timezone.now()
    run.report = {**run.report, "applied_new": created, "applied_updated": updated}
    run.save()
    return created, updated
