"""
Reads the catalogue spreadsheet (CSV or XLSX) into plain rows.

Column headings are matched by name, so column order does not matter and extra
columns (Order, Est. Time, Stock, Stores Location, the supplier block) are
ignored. The section for each row comes from, in order of preference:
a Section column, a section heading row above it, or the worksheet name.
"""

import csv
import io
import re
from dataclasses import asdict, dataclass, field
from decimal import Decimal, InvalidOperation

from .models import Section

FIELD_ALIASES = {
    "part_no": {"part no", "part number", "part", "code", "store code", "stores code", "item code", "product code"},
    "name": {"description", "item description", "name", "item", "item name", "material", "product", "product description"},
    "trade_price": {"trade price", "trade", "cost price", "cost", "unit cost", "trade cost", "price"},
    "sell_price": {"sell price", "sell", "selling price", "sale price"},
    "unit": {"unit", "uom", "unit of measure", "units", "unit of issue"},
    "section": {"section", "category", "department"},
    "flammable": {"flammable"},
}

SECTION_ALIASES = {
    Section.ADHESIVES: {"adhesive", "adhesives"},
    Section.BATHROOMS_KITCHENS: {"bathroomsandkitchens", "bathroomandkitchen", "bathroomsandkitchen",
                                "kitchensandbathrooms", "kitchenandbathroom", "bandk", "bathrooms", "kitchens"},
    Section.ELECTRICAL: {"electrical", "electric", "electrics"},
    Section.FLAMMABLE: {"flammableitems", "flammable", "flammables", "flammableitem"},
    Section.FLOORING: {"flooring", "floor", "floors"},
    Section.JOINERY: {"joinery"},
    Section.PLASTERING: {"plastering", "plaster"},
    Section.PLUMBING: {"plumbing"},
    Section.MISC: {"misc", "miscellaneous", "sundries"},
}


class SpreadsheetError(Exception):
    """The file could not be read as a catalogue."""


@dataclass
class SourceRow:
    sheet: str
    row_number: int
    part_no: str = ""
    name: str = ""
    trade_price: str | None = None  # Decimal as string (JSON friendly)
    sell_price: str | None = None
    unit: str = ""
    section: str | None = None  # Section key
    section_text: str = ""  # what the spreadsheet said, for the report
    flammable: bool = False
    price_text: str = ""  # raw price cell when it could not be read

    def to_dict(self):
        return asdict(self)

    @classmethod
    def from_dict(cls, data):
        return cls(**data)


@dataclass
class ReadResult:
    rows: list = field(default_factory=list)
    ignored_columns: list = field(default_factory=list)
    skipped: list = field(default_factory=list)  # {"where", "reason"}
    unknown_sections: list = field(default_factory=list)


def _norm_header(value):
    text = str(value or "").strip().lower()
    text = re.sub(r"[.:£()#/]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def match_section(text):
    """Return the Section key for a heading such as 'Bathrooms & Kitchens'."""
    if not text:
        return None
    norm = re.sub(r"[^a-z]", "", str(text).lower().replace("&", "and"))
    for key, aliases in SECTION_ALIASES.items():
        if norm in aliases:
            return key
    return None


def parse_price(value):
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value)).quantize(Decimal("0.0001"))
    text = str(value).replace("£", "").replace(",", "").replace("Â", "").strip()
    try:
        return Decimal(text).quantize(Decimal("0.0001"))
    except InvalidOperation:
        return None


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _map_columns(header_cells):
    mapping, ignored = {}, []
    for index, raw in enumerate(header_cells):
        norm = _norm_header(raw)
        if not norm:
            continue
        target = None
        if not norm.startswith("supplier"):
            for field_name, aliases in FIELD_ALIASES.items():
                if norm in aliases and field_name not in mapping:
                    target = field_name
                    break
        if target:
            mapping[target] = index
        else:
            ignored.append(str(raw).strip())
    return mapping, ignored


def _find_header(rows, limit=15):
    for position, cells in enumerate(rows[:limit]):
        normed = {_norm_header(c) for c in cells}
        if normed & FIELD_ALIASES["part_no"] and normed & FIELD_ALIASES["name"]:
            return position
    return None


def _read_table(sheet_name, rows, result):
    header_pos = _find_header(rows)
    if header_pos is None:
        if any(any(_cell_text(c) for c in r) for r in rows):
            result.skipped.append({"where": f"Sheet '{sheet_name}'",
                                   "reason": "No heading row with 'Part No.' and 'Description' columns found."})
        return
    mapping, ignored = _map_columns(rows[header_pos])
    for col in ignored:
        if col not in result.ignored_columns:
            result.ignored_columns.append(col)

    current_section = match_section(sheet_name)
    current_section_text = sheet_name if current_section else ""

    def get(cells, field_name):
        index = mapping.get(field_name)
        return cells[index] if index is not None and index < len(cells) else None

    for offset, cells in enumerate(rows[header_pos + 1:], start=header_pos + 2):
        texts = [_cell_text(c) for c in cells]
        if not any(texts):
            continue
        part_no = _cell_text(get(cells, "part_no")).upper()
        name = _cell_text(get(cells, "name"))
        price_cell = get(cells, "trade_price")

        # A section heading row: one bit of text that names a section, no price.
        filled = [t for t in texts if t]
        if len(filled) == 1 and parse_price(price_cell) is None and match_section(filled[0]):
            current_section, current_section_text = match_section(filled[0]), filled[0]
            continue
        if not part_no and not name:
            result.skipped.append({"where": f"{sheet_name} row {offset}", "reason": "No Part No. or name."})
            continue

        section_text = _cell_text(get(cells, "section"))
        section = match_section(section_text) if section_text else current_section
        if section_text and section is None and section_text not in result.unknown_sections:
            result.unknown_sections.append(section_text)
        trade = parse_price(price_cell)
        sell = parse_price(get(cells, "sell_price"))
        flammable_cell = _cell_text(get(cells, "flammable")).lower()

        result.rows.append(SourceRow(
            sheet=sheet_name,
            row_number=offset,
            part_no=part_no,
            name=name,
            trade_price=str(trade) if trade is not None else None,
            sell_price=str(sell) if sell is not None else None,
            unit=_cell_text(get(cells, "unit")),
            section=section,
            section_text=section_text or current_section_text,
            flammable=flammable_cell in ("y", "yes", "true", "1", "x"),
            price_text="" if trade is not None else _cell_text(price_cell),
        ))


def _decode_csv(data):
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def read_spreadsheet(filename, data):
    """Read a CSV or XLSX file (bytes) and return a ReadResult."""
    result = ReadResult()
    lower = filename.lower()
    if lower.endswith(".csv"):
        text = _decode_csv(data)
        rows = list(csv.reader(io.StringIO(text)))
        sheet = re.sub(r"\.csv$", "", filename.rsplit("/", 1)[-1], flags=re.I)
        _read_table(sheet, rows, result)
    elif lower.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - any parse failure is the same to the user
            raise SpreadsheetError(f"The file could not be opened as an Excel workbook ({exc}).") from exc
        for worksheet in workbook.worksheets:
            rows = [list(r) for r in worksheet.iter_rows(values_only=True)]
            _read_table(worksheet.title, rows, result)
        workbook.close()
    else:
        raise SpreadsheetError("Please upload a .csv or .xlsx file.")
    if not result.rows:
        raise SpreadsheetError(
            "No items found. The file needs a heading row with at least 'Part No.' and 'Description' columns."
        )
    return result
