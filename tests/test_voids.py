import io
from datetime import date

import openpyxl
import pytest
from django.urls import reverse
from openpyxl.styles import Font, PatternFill

from accounts.models import User
from voids import importer
from voids.models import Supervisor, Trade, TradeTask, Void

TODAY = date(2026, 10, 12)
GREEN = PatternFill("solid", fgColor="FF00B050")
ORANGE = PatternFill("solid", fgColor="FFFFC000")
YELLOW = PatternFill("solid", fgColor="FFFFFF00")
RED = Font(color="FFFF0000")


def tracker(rows, sheet="Voids", headings=None):
    """A small workbook laid out like the tracker: colour key, notes row, heading row."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws["A1"] = "ECC VOID TRACKER"
    for col, (label, fill) in zip("HIL", [("Works Completed", GREEN), ("Works in Progress", ORANGE),
                                          ("Snagging Issued", YELLOW)]):
        ws[f"{col}1"] = label
        ws[f"{col}1"].fill = fill
    ws.append([])
    ws.append(headings or ["Post Code", "Supervisor", "Void Area", "SIMPRO Number", "Expected Completion",
                           "Address", "Access", "Joinery", "Plumbing", "Electrical", "Painting"])
    for r in rows:
        values = [v[0] if isinstance(v, tuple) else v for v in r]
        ws.append(values)
        for i, v in enumerate(r):
            if isinstance(v, tuple):
                cell = ws.cell(ws.max_row, i + 1)
                for style in v[1:]:
                    if isinstance(style, PatternFill):
                        cell.fill = style
                    else:
                        cell.font = style
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


ROWS = [
    ["EH16 6PR", "Stee L", "SE", 7968, "Plumbing line req", "38 Balmwell Terrace", "Master Key",
     ("P Works 23/09", ORANGE, RED), ("Ben 23/09\nKyle 06/10", GREEN), "STC", ("Not needed",)],
    [None, None, None, None, None, "NEW SYSTEM"],
    ["EH4 4SH", "Andrew", "NW ", 11157, date(2026, 10, 2), "4/6 Muirhouse View", "",
     ("Ryan 22/9", GREEN, RED), ("Kyle", GREEN), ("Elite 15/09 - Rewire FF Done", YELLOW), None],
]


@pytest.fixture
def voids_user(make_user):
    return make_user("sup@bfsuk.org", "Sam Supervisor", role=User.Role.ORDERING)


@pytest.fixture
def voids_client(client, voids_user):
    voids_user.can_use_voids = True
    voids_user.save()
    client.force_login(voids_user)
    return client


def test_parse_reads_text_colours_and_skips_dividers(db):
    rows, report = importer.parse_workbook(tracker(ROWS), today=TODAY)
    assert [r["address"] for r in rows] == ["38 Balmwell Terrace", "4/6 Muirhouse View"]
    first = rows[0]
    assert first["expected_note"] == "Plumbing line req"
    joinery = first["tasks"]["Joinery"]
    assert joinery["status"] == TradeTask.Status.IN_PROGRESS and joinery["claimed"]
    assert joinery["assignee"] == "P Works" and joinery["date"] == "2026-09-23"
    plumbing = first["tasks"]["Plumbing"]
    assert plumbing["status"] == TradeTask.Status.COMPLETE and not plumbing["claimed"]
    assert plumbing["notes"] == "Ben 23/09\nKyle 06/10" and plumbing["date"] == "2026-10-06"
    assert first["tasks"]["Electrical"]["status"] == TradeTask.Status.STC
    assert first["tasks"]["Painting"]["status"] == TradeTask.Status.NOT_REQUIRED
    second = rows[1]
    assert second["expected"] == "2026-10-02" and second["area"] == "NW"
    assert second["tasks"]["Electrical"]["status"] == TradeTask.Status.SNAGGING
    assert second["electrical_flag"]


def test_columns_found_by_heading_not_position(db):
    headings = ["Address", "SIMPRO Number", "Painting", "Supervisor", "Mystery"]
    rows, report = importer.parse_workbook(
        tracker([["1 Test Street", 555, ("Decor8 01/10", GREEN), "Martin", "x"]], headings=headings), today=TODAY)
    assert rows[0]["simpro"] == "555"
    assert rows[0]["tasks"]["Painting"]["status"] == TradeTask.Status.COMPLETE
    assert any("Mystery" in p for p in report["problems"])


def test_apply_creates_then_updates_without_duplicates(db, admin_user):
    run = importer.preview_upload("t.xlsx", tracker(ROWS), admin_user, today=TODAY)
    assert run.report["new"] == 2
    assert importer.apply_run(run, admin_user) == (2, 0)
    v = Void.objects.get(simpro="7968")
    assert v.job_type.name == "Void"
    assert v.supervisor.name == "Stee L" and v.area.code == "SE"
    assert v.status == Void.Status.IN_PROGRESS
    muir = Void.objects.get(simpro="11157")
    assert muir.supervisor.name == "Andrew F"  # matched by import alias
    assert muir.status == Void.Status.SNAGGING  # all complete except a snagging trade
    run2 = importer.preview_upload("t.xlsx", tracker(ROWS), admin_user, today=TODAY)
    assert run2.report["new"] == 0 and run2.report["updated"] == 2
    importer.apply_run(run2, admin_user)
    assert Void.objects.count() == 2


def test_postcode_sets_patch_supervisor_for_new_void(voids_client):
    from voids.models import JobType
    resp = voids_client.post(reverse("voids:new"), {
        "job_type": JobType.objects.get(name="Temp").pk, "address": "9 New Street", "postcode": "eh11 2aa",
        "status": "new"})
    assert resp.status_code == 302
    v = Void.objects.get(address="9 New Street")
    assert v.supervisor.name == "Martin T"
    assert v.status == Void.Status.NEW


def test_grid_cell_edit_updates_status(voids_client):
    from voids.models import JobType
    v = Void.objects.create(job_type=JobType.objects.first(), address="1 Grid Road", status=Void.Status.STC)
    trade = Trade.objects.get(name="Joinery")
    url = reverse("voids:task_cell", args=[v.pk, trade.pk])
    assert b"cellform" in voids_client.get(url, HTTP_HX_REQUEST="true").content
    resp = voids_client.post(url, {"status": "complete", "assignee": "Kev", "date": "2026-10-01"},
                             HTTP_HX_REQUEST="true")
    assert b"t-complete" in resp.content
    v.refresh_from_db()
    assert v.status == Void.Status.COMPLETE


def test_filters_by_job_type_and_electrical_supervisor(voids_client):
    from voids.models import JobType
    temp, void_type = JobType.objects.get(name="Temp"), JobType.objects.get(name="Void")
    Void.objects.create(job_type=temp, address="Temp Flat", status=Void.Status.STC)
    sparky = Void.objects.create(job_type=void_type, address="Sparky Flat", status=Void.Status.STC,
                                 electrical_flag=True)
    resp = voids_client.get(reverse("voids:list"), {"job": temp.pk})
    assert b"Temp Flat" in resp.content and b"Sparky Flat" not in resp.content
    jamie = Supervisor.objects.get(name="Jamie H")
    resp = voids_client.get(reverse("voids:list"), {"supervisor": jamie.pk})
    assert b"Sparky Flat" in resp.content and b"Temp Flat" not in resp.content
    assert sparky in Void.objects.all()


def test_module_access(client, make_user):
    orderer = make_user("o@bfsuk.org", "Olly Orders")
    client.force_login(orderer)
    assert client.get(reverse("voids:list")).status_code == 403
    assert client.get(reverse("orders:list")).status_code == 200
    voids_only = make_user("v@bfsuk.org", "Vic Voids")
    voids_only.can_use_orders, voids_only.can_use_voids = False, True
    voids_only.save()
    client.force_login(voids_only)
    assert client.get(reverse("orders:list")).status_code == 403
    assert client.get(reverse("voids:list")).status_code == 200
    assert client.get(reverse("home"))["Location"] == reverse("voids:list")


def test_admin_switches_module_access(admin_client, make_user):
    u = make_user("x@bfsuk.org", "Xan")
    admin_client.post(reverse("accounts:user_action", args=[u.pk]), {"action": "grant_voids"})
    u.refresh_from_db()
    assert u.can_use_voids


def test_export_and_import_pages(admin_client):
    assert admin_client.get(reverse("voids:export")).status_code == 200
    assert admin_client.get(reverse("voids:import")).status_code == 200


def test_tiles_by_area_and_tile_edit(voids_client):
    from voids.models import Area, JobType
    ne, sw = Area.objects.get(code="NE"), Area.objects.get(code="SW")
    a = Void.objects.create(job_type=JobType.objects.first(), address="North Flat", area=ne, status=Void.Status.STC)
    Void.objects.create(job_type=JobType.objects.first(), address="South Flat", area=sw, status=Void.Status.STC)
    resp = voids_client.get(reverse("voids:list"), {"view": "tiles", "area": ne.pk})
    assert b"North Flat" in resp.content and b"South Flat" not in resp.content
    assert b'class="tiles"' in resp.content
    trade = Trade.objects.get(name="Plumbing")
    url = reverse("voids:task_cell", args=[a.pk, trade.pk]) + "?in=tile"
    resp = voids_client.post(url, {"status": "in_progress", "assignee": "Ben"}, HTTP_HX_REQUEST="true")
    assert b'class="tile' in resp.content and b"dot t-in_progress" in resp.content
    a.refresh_from_db()
    assert a.status == Void.Status.IN_PROGRESS


def test_diary_update_visit_and_stale_flags(voids_client, voids_user, settings):
    from datetime import timedelta

    from django.utils import timezone

    from voids.models import JobType, VoidUpdate
    jt = JobType.objects.first()
    fresh = Void.objects.create(job_type=jt, address="Fresh Flat", status=Void.Status.NEW,
                                key_info="Key safe 4321")
    old = Void.objects.create(job_type=jt, address="Old Flat", status=Void.Status.IN_PROGRESS,
                              last_activity=timezone.now() - timedelta(days=12))
    mid = Void.objects.create(job_type=jt, address="Mid Flat", status=Void.Status.STC,
                              last_activity=timezone.now() - timedelta(days=8))
    Void.objects.create(job_type=jt, address="Held Flat", status=Void.Status.STC, on_hold=True,
                        last_activity=timezone.now() - timedelta(days=30))
    assert old.stale_level == "red" and mid.stale_level == "amber" and fresh.stale_level == ""

    resp = voids_client.get(reverse("voids:list"), {"flag": "stale", "sort": "stale"})
    body = resp.content.decode()
    assert "Held Flat" not in body and "Fresh Flat" not in body
    assert body.index("Old Flat") < body.index("Mid Flat")  # oldest first
    assert "No update · 12 days" in body

    resp = voids_client.get(reverse("voids:list"), {"flag": "", "sort": ""})
    assert b"Key safe 4321" in resp.content

    # A site visit from a tile clears "needs visit", sets the inspection date and resets the clock.
    resp = voids_client.post(reverse("voids:add_update", args=[old.pk]),
                             {"kind": "visit", "text": "Visited, plasterer booked"}, HTTP_HX_REQUEST="true")
    assert b"Visited, plasterer booked" in resp.content
    old.refresh_from_db()
    assert old.stale_level == "" and old.last_inspection == timezone.localdate()
    voids_client.post(reverse("voids:add_update", args=[fresh.pk]), {"kind": "visit", "text": "First visit"})
    fresh.refresh_from_db()
    assert fresh.status == Void.Status.STC

    # Trade changes are written to the diary automatically, with the optional comment.
    trade = Trade.objects.get(name="Joinery")
    voids_client.post(reverse("voids:task_cell", args=[mid.pk, trade.pk]),
                      {"status": "complete", "assignee": "Kev", "comment": "all done"})
    texts = list(VoidUpdate.objects.filter(void=mid).values_list("text", flat=True))
    assert any(t.startswith("Joinery → Complete (Kev)") and t.endswith("all done") for t in texts)
    mid.refresh_from_db()
    assert mid.stale_level == ""

    upd = VoidUpdate.objects.get(void=old, kind="visit")
    voids_client.post(reverse("voids:edit_update", args=[upd.pk]), {"kind": "visit", "text": "Fixed typo"})
    upd.refresh_from_db()
    assert upd.text == "Fixed typo" and upd.edited_at
    VoidUpdate.objects.filter(pk=upd.pk).update(created_at=timezone.now() - timedelta(minutes=20))
    voids_client.post(reverse("voids:edit_update", args=[upd.pk]), {"kind": "visit", "text": "Too late"})
    upd.refresh_from_db()
    assert upd.text == "Fixed typo"


def test_import_does_not_count_as_activity(db, admin_user):
    from datetime import timedelta

    from django.utils import timezone
    run = importer.preview_upload("t.xlsx", tracker(ROWS), admin_user, today=TODAY)
    importer.apply_run(run, admin_user)
    Void.objects.update(last_activity=timezone.now() - timedelta(days=20))
    run = importer.preview_upload("t.xlsx", tracker(ROWS), admin_user, today=TODAY)
    importer.apply_run(run, admin_user)
    assert all(v.stale_level == "red" for v in Void.objects.all())


def test_add_trade_starts_stc_and_remove(voids_client):
    from voids.models import JobType
    v = Void.objects.create(job_type=JobType.objects.first(), address="Trade Flat", status=Void.Status.STC)
    plumbing = Trade.objects.get(name="Plumbing")
    resp = voids_client.post(reverse("voids:add_trade", args=[v.pk]), {"trade": plumbing.pk}, HTTP_HX_REQUEST="true")
    assert b"Trades on this job" in resp.content
    task = TradeTask.objects.get(void=v, trade=plumbing)
    assert task.status == TradeTask.Status.STC
    voids_client.post(reverse("voids:task_cell", args=[v.pk, plumbing.pk]), {"remove": "1"})
    assert not TradeTask.objects.filter(void=v).exists()


def test_issue_update_flags_job_until_cleared(voids_client):
    from voids.models import JobType
    v = Void.objects.create(job_type=JobType.objects.first(), address="Issue Flat", status=Void.Status.STC)
    resp = voids_client.post(reverse("voids:add_update", args=[v.pk]),
                             {"kind": "issue", "text": "Leak under sink"}, HTTP_HX_REQUEST="true")
    assert b"ISSUE: Leak under sink" in resp.content
    v.refresh_from_db()
    assert v.issue_flag and v.issue_note == "Leak under sink"
    assert b"Issue Flat" in voids_client.get(reverse("voids:list"), {"flag": "issue"}).content
    voids_client.post(reverse("voids:clear_issue", args=[v.pk]))
    v.refresh_from_db()
    assert not v.issue_flag


def test_sort_by_job_number_and_electrical_works_filter(voids_client):
    from voids.models import JobType
    jt = JobType.objects.first()
    for simpro, addr in [("12000", "C Flat"), ("900", "A Flat"), ("10584/10995", "B Flat")]:
        Void.objects.create(job_type=jt, address=addr, simpro=simpro, status=Void.Status.STC)
    body = voids_client.get(reverse("voids:list"), {"sort": "", "flag": ""}).content.decode()
    assert body.index("A Flat") < body.index("B Flat") < body.index("C Flat")
    sparky = Void.objects.create(job_type=jt, address="Rewire Flat", status=Void.Status.STC)
    TradeTask.objects.create(void=sparky, trade=Trade.objects.get(name="Electrical"))
    body = voids_client.get(reverse("voids:list"), {"flag": "electrical"}).content.decode()
    assert "Rewire Flat" in body and "A Flat" not in body


def test_thresholds_are_editable(voids_client, admin_client):
    from datetime import timedelta

    from django.utils import timezone

    from voids.models import JobType, VoidsSettings
    v = Void.objects.create(job_type=JobType.objects.first(), address="Slow Flat", status=Void.Status.STC,
                            last_activity=timezone.now() - timedelta(days=4))
    assert v.stale_level == ""
    resp = admin_client.post(reverse("admin:voids_voidssettings_change", args=[VoidsSettings.get().pk]), {
        "stale_amber_days": 3, "stale_red_days": 6, "stale_filter_days": 2, "handback_soon_days": 14})
    assert resp.status_code == 302
    v = Void.objects.get(pk=v.pk)
    assert v.stale_level == "amber"
    body = voids_client.get(reverse("voids:list"), {"flag": "stale"}).content.decode()
    assert "Slow Flat" in body and "No update in 2+ days" in body
    resp = admin_client.post(reverse("admin:voids_voidssettings_change", args=[1]), {
        "stale_amber_days": 6, "stale_red_days": 3, "stale_filter_days": 2, "handback_soon_days": 14})
    assert resp.status_code == 200  # red must be more than amber


def test_all_trades_complete_requests_final_check_then_completed(voids_client, voids_user, mailoutbox):
    from voids.models import JobType, VoidUpdate
    sup = Supervisor.objects.get(name="Andrew F")
    sup.email = "andrew@bfsuk.org"
    sup.save()
    v = Void.objects.create(job_type=JobType.objects.first(), address="Done Flat", supervisor=sup,
                            status=Void.Status.IN_PROGRESS, simpro="555")
    j, p = Trade.objects.get(name="Joinery"), Trade.objects.get(name="Plumbing")
    TradeTask.objects.create(void=v, trade=j, status="complete")
    TradeTask.objects.create(void=v, trade=p, status="in_progress")
    voids_client.post(reverse("voids:task_cell", args=[v.pk, p.pk]) + "?in=tile",
                      {"status": "complete", "autosave": "1"}, HTTP_HX_REQUEST="true")
    v.refresh_from_db()
    assert v.status == Void.Status.COMPLETE and v.final_check_requested_at
    assert len(mailoutbox) == 1 and mailoutbox[0].to == ["andrew@bfsuk.org"]
    assert "Final check needed: Done Flat" in mailoutbox[0].subject
    assert VoidUpdate.objects.filter(void=v, text__contains="final check requested").exists()
    # Not in the Completed list until the final check is confirmed.
    assert b"Done Flat" not in voids_client.get(reverse("voids:list"), {"status": "handed_back"}).content
    resp = voids_client.post(reverse("voids:final_check", args=[v.pk]), HTTP_HX_REQUEST="true")
    assert b"final checked by" in resp.content
    v.refresh_from_db()
    assert v.status == Void.Status.HANDED_BACK and v.final_checked_by == voids_user
    assert b"Done Flat" in voids_client.get(reverse("voids:list"), {"status": "handed_back"}).content
    assert b"Done Flat" not in voids_client.get(reverse("voids:list"), {"status": ""}).content
    voids_client.post(reverse("voids:final_check", args=[v.pk]), {"reopen": "1"})
    v.refresh_from_db()
    assert v.status != Void.Status.HANDED_BACK


def test_autosave_keeps_editor_open_and_merges_diary_lines(voids_client, voids_user):
    from voids.models import JobType, VoidUpdate
    v = Void.objects.create(job_type=JobType.objects.first(), address="Auto Flat", status=Void.Status.STC)
    t = Trade.objects.get(name="Joinery")
    TradeTask.objects.create(void=v, trade=t)
    url = reverse("voids:task_cell", args=[v.pk, t.pk]) + "?in=tile"
    resp = voids_client.post(url, {"status": "in_progress", "assignee": "Kev", "autosave": "1"}, HTTP_HX_REQUEST="true")
    assert b"autosave" in resp.content and b"Saved" in resp.content  # still editing
    voids_client.post(url, {"status": "in_progress", "assignee": "Kev", "notes": "first fix", "autosave": "1"},
                      HTTP_HX_REQUEST="true")
    lines = list(VoidUpdate.objects.filter(void=v, kind="trade").values_list("text", flat=True))
    assert lines == ["Joinery → In progress (Kev)"]
    task = TradeTask.objects.get(void=v, trade=t)
    assert task.updated_by == voids_user and task.assignee == "Kev"


def test_one_click_claim_and_claim_all(voids_client):
    from voids.models import JobType
    v = Void.objects.create(job_type=JobType.objects.first(), address="Claim Flat", status=Void.Status.IN_PROGRESS)
    j, p, e = (Trade.objects.get(name=n) for n in ("Joinery", "Plumbing", "Electrical"))
    TradeTask.objects.create(void=v, trade=j, status="complete")
    TradeTask.objects.create(void=v, trade=p, status="complete")
    TradeTask.objects.create(void=v, trade=e, status="in_progress")
    resp = voids_client.post(reverse("voids:toggle_claim", args=[v.pk, e.pk]), HTTP_HX_REQUEST="true")
    assert b"Claimed" in resp.content and TradeTask.objects.get(void=v, trade=e).claimed
    voids_client.post(reverse("voids:claim_complete", args=[v.pk]), HTTP_HX_REQUEST="true")
    assert TradeTask.objects.filter(void=v, claimed=True).count() == 3


def test_supervisor_buttons_and_own_voids_by_default(client, make_user):
    from voids.models import JobType
    jt_void, jt_temp = JobType.objects.get(name="Void"), JobType.objects.get(name="Temp")
    stee, martin, jamie = (Supervisor.objects.get(name=n) for n in ("Stee L", "Martin T", "Jamie H"))
    Void.objects.create(job_type=jt_void, address="Stee Flat", supervisor=stee, status="stc")
    m = Void.objects.create(job_type=jt_temp, address="Martin Temp Flat", supervisor=martin, status="stc")
    TradeTask.objects.create(void=m, trade=Trade.objects.get(name="Electrical"))
    user = make_user("stee@bfsuk.org", "Stee Lee")
    user.can_use_voids = True
    user.save()
    stee.user = user
    stee.save()
    client.force_login(user)
    body = client.get(reverse("voids:list")).content.decode()
    assert "Stee Flat" in body and "Martin Temp Flat" not in body  # starts on own voids
    assert "or by supervisor" in body and "Jamie H (Electrical)" in body
    body = client.get(reverse("voids:list"), {"supervisor": jamie.pk}).content.decode()
    assert "Martin Temp Flat" in body and "Stee Flat" not in body
    # Electrical button shows electrical voids of every job type, with electrical prominent.
    body = client.get(reverse("voids:list"), {"supervisor": "", "job": jt_void.pk, "flag": "electrical"}).content.decode()
    assert "Martin Temp Flat" in body and "elec-bar" in body


def test_sign_off_only_when_all_trades_complete_and_button_on_tile(voids_client):
    from voids import services
    from voids.models import JobType
    v = Void.objects.create(job_type=JobType.objects.first(), address="Sign Flat", status=Void.Status.IN_PROGRESS)
    j, p = Trade.objects.get(name="Joinery"), Trade.objects.get(name="Plumbing")
    TradeTask.objects.create(void=v, trade=j, status="complete")
    pt = TradeTask.objects.create(void=v, trade=p, status="in_progress")
    voids_client.post(reverse("voids:final_check", args=[v.pk]))
    v.refresh_from_db()
    assert v.status != Void.Status.HANDED_BACK  # plumbing still in progress
    pt.delete()  # removed from the job
    services.refresh_status(v)
    resp = voids_client.get(reverse("voids:list"), {"q": "Sign Flat", "status": ""})
    assert b"Confirm final check" in resp.content
    voids_client.post(reverse("voids:final_check", args=[v.pk]))
    v.refresh_from_db()
    assert v.status == Void.Status.HANDED_BACK
    assert f'id="tile-{v.pk}"'.encode() not in voids_client.get(reverse("voids:list"), {"q": "", "status": ""}).content


def test_days_left_auto_electrical_and_auto_inspection(voids_client):
    from datetime import timedelta

    from django.utils import timezone

    from voids.models import JobType
    today = timezone.localdate()
    v = Void.objects.create(job_type=JobType.objects.first(), address="Count Flat", status=Void.Status.STC,
                            handback_date=today + timedelta(days=5))
    assert v.handback_label == "5 days left"
    v.handback_date = today - timedelta(days=1)
    assert v.handback_label == "1 day overdue"
    e = Trade.objects.get(name="Electrical")
    voids_client.post(reverse("voids:add_trade", args=[v.pk]), {"trade": e.pk})
    v.refresh_from_db()
    assert v.electrical_flag
    voids_client.post(reverse("voids:task_cell", args=[v.pk, e.pk]), {"status": "in_progress", "assignee": "Elite"})
    v.refresh_from_db()
    assert v.last_inspection == today
    voids_client.post(reverse("voids:task_cell", args=[v.pk, e.pk]), {"remove": "1"})
    v.refresh_from_db()
    assert not v.electrical_flag
    form_html = voids_client.get(reverse("voids:new")).content.decode()
    for gone in ("keys_received", "status_fixed", "electrical_flag", "last_inspection"):
        assert f'name="{gone}"' not in form_html


def test_nightly_email_puts_snagging_first(db, mailoutbox):
    from voids import services
    from voids.digest import send_all
    from voids.models import JobType
    sup = Supervisor.objects.get(name="Stee L")
    sup.email = "stee@bfsuk.org"
    sup.save()
    jt = JobType.objects.first()
    snag = Void.objects.create(job_type=jt, address="Snag Flat", supervisor=sup, status=Void.Status.IN_PROGRESS)
    TradeTask.objects.create(void=snag, trade=Trade.objects.get(name="Painting"), status="snagging")
    services.refresh_status(snag)
    snag.refresh_from_db()
    assert snag.status == Void.Status.SNAGGING and snag.snagging_since
    Void.objects.create(job_type=jt, address="Visit Flat", supervisor=sup, status=Void.Status.NEW)
    report = send_all()
    assert len(mailoutbox) == 1
    msg = mailoutbox[0]
    assert msg.to == ["stee@bfsuk.org"] and "1 snagging" in msg.subject
    html = msg.alternatives[0][0]
    assert html.index("SNAGGING") < html.index("Visit Flat")
    assert "Snagging: Painting" in html
    assert any(r[0] == sup and r[3] for r in report)


def test_overview_counts_and_links(voids_client):
    from datetime import timedelta

    from django.utils import timezone

    from voids.models import JobType
    today = timezone.localdate()
    jt = JobType.objects.first()
    stee = Supervisor.objects.get(name="Stee L")
    Void.objects.create(job_type=jt, address="Late Flat", supervisor=stee, status="in_progress",
                        handback_date=today - timedelta(days=2), issue_flag=True)
    v = Void.objects.create(job_type=jt, address="Soon Flat", supervisor=stee, status="snagging",
                            handback_date=today + timedelta(days=3))
    TradeTask.objects.create(void=v, trade=Trade.objects.get(name="Painting"), status="stc")
    resp = voids_client.get(reverse("voids:overview"))
    assert resp.status_code == 200
    ctx = resp.context
    assert ctx["totals"]["open"] == 2 and ctx["totals"]["overdue"] == 1 and ctx["totals"]["issues"] == 1
    assert ctx["totals"]["snagging"] == 1 and ctx["totals"]["due_soon"] == 1
    stee_row = next(r for r in ctx["sup_rows"] if r["label"] == "Stee L")
    overdue_cell = next(c for c in stee_row["cells"] if c["key"] == "overdue")
    assert overdue_cell["n"] == 1 and f"supervisor={stee.pk}" in overdue_cell["url"] and "due=overdue" in overdue_cell["url"]
    assert ctx["trades"][0]["trade"].name == "Painting"
    body = voids_client.get(overdue_cell["url"]).content.decode()
    assert "Late Flat" in body and "Soon Flat" not in body


def test_electrical_snagging_goes_to_jamie_not_patch_supervisor(db, mailoutbox):
    from voids import services
    from voids.digest import send_all
    from voids.models import JobType
    stee, jamie = Supervisor.objects.get(name="Stee L"), Supervisor.objects.get(name="Jamie H")
    stee.email, jamie.email = "stee@bfsuk.org", "jamie@bfsuk.org"
    stee.save()
    jamie.save()
    v = Void.objects.create(job_type=JobType.objects.first(), address="Spark Snag Flat", supervisor=stee,
                            status="in_progress")
    TradeTask.objects.create(void=v, trade=Trade.objects.get(name="Electrical"), status="snagging")
    TradeTask.objects.create(void=v, trade=Trade.objects.get(name="Joinery"), status="complete")
    services.refresh_status(v)
    send_all()
    by_to = {m.to[0]: m for m in mailoutbox}
    assert "Snagging: Electrical" in by_to["jamie@bfsuk.org"].alternatives[0][0]
    assert "snagging" in by_to["jamie@bfsuk.org"].subject
    stee_mail = by_to.get("stee@bfsuk.org")
    assert stee_mail is None or "Snagging: Electrical" not in stee_mail.alternatives[0][0]


def test_status_buttons_set_dates_and_area_or_supervisor(voids_client):
    from django.utils import timezone

    from voids.models import Area, JobType
    v = Void.objects.create(job_type=JobType.objects.first(), address="Tap Flat", status="stc")
    t = Trade.objects.get(name="Joinery")
    TradeTask.objects.create(void=v, trade=t, assignee="Kev")
    resp = voids_client.post(reverse("voids:set_task_status", args=[v.pk, t.pk]), {"status": "in_progress"},
                             HTTP_HX_REQUEST="true")
    assert b"st-btn t-in_progress on" in resp.content
    task = TradeTask.objects.get(void=v, trade=t)
    assert task.status == "in_progress" and task.date == timezone.localdate()
    v.refresh_from_db()
    assert v.status == Void.Status.IN_PROGRESS and v.last_inspection == timezone.localdate()
    # Area and supervisor are either/or: picking a supervisor ignores the area.
    stee = Supervisor.objects.get(name="Stee L")
    se = Area.objects.get(code="SE")
    Void.objects.create(job_type=JobType.objects.first(), address="Stee NW Flat", supervisor=stee,
                        area=Area.objects.get(code="NW"), status="stc")
    body = voids_client.get(reverse("voids:list"), {"area": se.pk, "supervisor": stee.pk}).content.decode()
    assert "Stee NW Flat" in body


def test_no_booked_status():
    assert "booked" not in TradeTask.Status.values
    from voids.views import QUICK_STATUSES
    assert [s for s, _ in QUICK_STATUSES] == ["stc", "in_progress", "snagging", "complete"]


def test_in_progress_asks_who(voids_client):
    from voids.models import JobType
    v = Void.objects.create(job_type=JobType.objects.first(), address="Who Flat", status="stc")
    t = Trade.objects.get(name="Plumbing")
    TradeTask.objects.create(void=v, trade=t)
    # Button with nobody assigned: opens the editor asking who, nothing saved yet.
    resp = voids_client.post(reverse("voids:set_task_status", args=[v.pk, t.pk]), {"status": "in_progress"},
                             HTTP_HX_REQUEST="true")
    assert b"Who is doing the plumbing?" in resp.content and b"autofocus" in resp.content
    from django.utils import timezone
    assert timezone.localdate().isoformat().encode() in resp.content  # date filled in, still editable
    assert TradeTask.objects.get(void=v, trade=t).status == "stc"
    url = reverse("voids:task_cell", args=[v.pk, t.pk]) + "?in=tile"
    resp = voids_client.post(url, {"status": "in_progress", "assignee": "", "autosave": "1"}, HTTP_HX_REQUEST="true")
    assert b"Who is doing it?" in resp.content
    assert TradeTask.objects.get(void=v, trade=t).status == "stc"
    voids_client.post(url, {"status": "in_progress", "assignee": "Kyle", "autosave": "1"}, HTTP_HX_REQUEST="true")
    task = TradeTask.objects.get(void=v, trade=t)
    assert task.status == "in_progress" and task.assignee == "Kyle"
    # Once someone is assigned the button changes status straight away.
    voids_client.post(reverse("voids:set_task_status", args=[v.pk, t.pk]), {"status": "snagging"})
    voids_client.post(reverse("voids:set_task_status", args=[v.pk, t.pk]), {"status": "in_progress"})
    assert TradeTask.objects.get(void=v, trade=t).status == "in_progress"
