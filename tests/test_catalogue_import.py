"""Catalogue import: temporary codes, flags, report, re-import and admin editing."""

import io
from decimal import Decimal

import pytest
from django.urls import reverse
from openpyxl import Workbook

from audit.models import AuditEntry
from catalogue import importer
from catalogue.models import ImportRun, Item, MeasureType, Section

pytestmark = pytest.mark.django_db

HEADERS = ["Order", "Part No.", "Description", "Unit", "Trade Price", "Sell Price", "Est. Time", "Stock",
           "Stores Location", "Supplier", "Supplier Price"]


def row(code, name, price, unit="each", sell=None):
    return [1, code, name, unit, price, sell if sell is not None else price, "5 min", 12, "Bay 3", "Acme", 1.0]


def xlsx(sheets):
    """sheets: {sheet name: [rows]} -> xlsx bytes, each sheet with the standard heading row."""
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        ws.append(HEADERS)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run_import(sheets, actor=None, apply=True, filename="catalogue.xlsx"):
    run = importer.preview_upload(filename, xlsx(sheets), actor)
    if apply:
        report = importer.apply_run(run, actor)
        run.refresh_from_db()
        return run, report
    return run, run.report


def item(code):
    return Item.objects.get(part_no=code)


# --- Duplicate and placeholder codes ---------------------------------------

def test_duplicate_part_numbers_get_temporary_suffixes_and_do_not_block_import():
    _, report = run_import({"Electrical": [
        row("ELE118", "Twin socket", 4.10),
        row("ELE118", "Single socket", 3.20),
        row("ELE119", "Switch", 2.00),
    ], "Joinery": [
        row("JOI095", "Door stop", 1.00),
        row("JOI095", "Door wedge", 1.10),
        row("JOI095", "Door handle", 6.50),
    ]})
    assert Item.objects.count() == 6
    first = item("ELE118")
    assert first.catalogue_name == "Twin socket" and not first.code_to_confirm
    second = item("ELE118-B")
    assert second.catalogue_name == "Single socket" and second.code_to_confirm
    assert second.original_part_no == "ELE118"
    assert item("JOI095-B").code_to_confirm and item("JOI095-C").code_to_confirm
    assert not item("ELE119").code_to_confirm
    pairs = {(t["original"], t["new"], t["reason"]) for t in report["temp_codes"]}
    assert pairs == {
        ("ELE118", "ELE118-B", "duplicate Part No."),
        ("JOI095", "JOI095-B", "duplicate Part No."),
        ("JOI095", "JOI095-C", "duplicate Part No."),
    }


def test_placeholder_codes_all_get_temporary_codes():
    _, report = run_import({"Plumbing": [
        row("PLU0", "Mystery fitting", 1.00),
        row("PLU0", "Another mystery", 2.00),
        row("PLU001", "Real fitting", 3.00),
    ]})
    a, b = item("PLU0-A"), item("PLU0-B")
    assert a.catalogue_name == "Mystery fitting" and a.code_to_confirm
    assert b.catalogue_name == "Another mystery" and b.code_to_confirm
    assert not Item.objects.filter(part_no="PLU0").exists()
    assert {t["new"] for t in report["temp_codes"]} == {"PLU0-A", "PLU0-B"}
    assert all(t["reason"] == "placeholder code" for t in report["temp_codes"])


def test_reimport_keeps_temporary_codes_stable():
    sheets = {"Electrical": [row("ELE118", "Twin socket", 4.10), row("ELE118", "Single socket", 3.20)],
              "Plumbing": [row("PLU0", "Mystery", 1.0)]}
    run_import(sheets)
    _, report = run_import(sheets)
    assert Item.objects.count() == 3
    assert report["summary"]["additions"] == 0 and report["summary"]["unchanged"] == 3
    assert report["temp_codes"] == []


def test_corrected_temporary_code_survives_reimport():
    sheets = {"Electrical": [row("ELE118", "Twin socket", 4.10), row("ELE118", "Single socket", 3.20)]}
    run_import(sheets)
    corrected = item("ELE118-B")
    corrected.part_no = "ELE130"
    corrected.code_to_confirm = False
    corrected.save()

    _, report = run_import(sheets)  # spreadsheet still has the duplicate
    assert Item.objects.count() == 2
    assert item("ELE130").catalogue_name == "Single socket" and item("ELE130").is_active
    assert report["removals"] == [] and report["additions"] == []

    # Spreadsheet later fixed to use the real code: same item, no duplicate.
    _, report = run_import({"Electrical": [row("ELE118", "Twin socket", 4.10), row("ELE130", "Single socket", 3.20)]})
    assert Item.objects.count() == 2 and report["removals"] == [] and report["additions"] == []


# --- Flags and report ---------------------------------------------------------

def test_malformed_rows_imported_as_incomplete_and_not_orderable():
    _, report = run_import({"Plastering": [
        row("PLA001", "Bonding plaster 25kg", None),
        row("PLA002", "", 5.00),
        row("PLA003", "Multi finish 25kg", "TBC"),
        row("PLA004", "Board 2.4m", 9.00),
    ]})
    for code in ("PLA001", "PLA002", "PLA003"):
        assert item(code).is_incomplete and not item(code).is_orderable
    assert item("PLA004").is_orderable
    listed = {x["part_no"]: x["missing"] for x in report["incomplete"]}
    assert listed == {"PLA001": "trade price", "PLA002": "name", "PLA003": "trade price"}

    # Completing the item clears the flag.
    fixed = item("PLA001")
    fixed.trade_price = Decimal("7.50")
    fixed.save()
    assert fixed.is_orderable


def test_sell_below_trade_reported_and_trade_price_used():
    _, report = run_import({"Adhesives": [row("ADH001", "Grab adhesive", 5.00, sell=4.00),
                                          row("ADH002", "PVA 1L", 3.00, sell=4.50)]})
    assert [x["part_no"] for x in report["sell_below_trade"]] == ["ADH001"]
    assert item("ADH001").trade_price == Decimal("5.0000")


def test_character_repair_and_remaining_garbled_names():
    _, report = run_import({"Flooring": [
        row("V0010", "Safety vinyl 2m âˆ’ grey", 20.0, unit="m2"),
        row("V0011", "Underlay 15mÂ² roll", 30.0),
        row("V0012", "Tile Ã\x81 broken", 1.0),
    ]})
    assert item("V0010").catalogue_name == "Safety vinyl 2m - grey"
    assert item("V0011").display_name == "Underlay 15m² roll"
    assert item("V0012").name_check
    assert [x["part_no"] for x in report["garbled_names"]] == ["V0012"]
    assert report["summary"]["repaired_names"] >= 2


def test_ignored_columns_reported():
    _, report = run_import({"Adhesives": [row("ADH001", "Grab adhesive", 5.00)]})
    for col in ("Order", "Est. Time", "Stock", "Stores Location", "Supplier", "Supplier Price"):
        assert col in report["ignored_columns"]


def test_preview_changes_nothing_until_applied():
    run, report = run_import({"Adhesives": [row("ADH001", "Grab adhesive", 5.00)]}, apply=False)
    assert run.status == ImportRun.Status.PREVIEW
    assert report["summary"]["additions"] == 1
    assert Item.objects.count() == 0
    importer.apply_run(run, None)
    assert Item.objects.count() == 1
    with pytest.raises(ValueError):
        importer.apply_run(run, None)  # cannot be applied twice


# --- Re-import -------------------------------------------------------------

def test_reimport_updates_name_and_price_but_never_display_name_unit_or_pack_size():
    run_import({"Plumbing": [row("PLU010", "Pipe clip 15mm per 100 - BRAND X", 4.00)]})
    clip = item("PLU010")
    assert clip.measure_type == MeasureType.PACK and clip.pack_size == 100
    clip.display_name = "Pipe clip 15mm"
    clip.pack_size = 50
    clip.unit = "bag"
    clip.save()

    _, report = run_import({"Plumbing": [row("PLU010", "Pipe clip 15mm per 100 - BRAND Y", 4.40)]})
    clip.refresh_from_db()
    assert clip.display_name == "Pipe clip 15mm"
    assert clip.catalogue_name == "Pipe clip 15mm per 100 - BRAND Y"
    assert clip.trade_price == Decimal("4.4000")
    assert (clip.pack_size, clip.unit) == (50, "bag")
    changes = report["updates"][0]["changes"]
    assert changes["Trade price"] == {"from": "4.00", "to": "4.40"}
    assert AuditEntry.objects.filter(action="item.import_update", target_id=str(clip.pk)).exists()


def test_items_missing_from_file_are_deactivated_only_in_sections_covered():
    run_import({"Adhesives": [row("ADH001", "Grab", 5.0), row("ADH002", "PVA", 3.0)],
                "Joinery": [row("JOI001", "Hinge", 2.0)]})
    run, report = run_import({"Adhesives": [row("ADH001", "Grab", 5.0)]}, apply=False)
    assert [x["part_no"] for x in report["removals"]] == ["ADH002"]  # Joinery not in file: untouched
    importer.apply_run(run, None)
    assert not item("ADH002").is_active
    assert item("JOI001").is_active
    # Comes back if re-imported.
    run_import({"Adhesives": [row("ADH001", "Grab", 5.0), row("ADH002", "PVA", 3.0)]})
    assert item("ADH002").is_active


def test_manual_items_are_never_deactivated_by_import():
    Item.objects.create(part_no="BFS99", display_name="Admin added", section_id=Section.MISC, trade_price=1)
    run_import({"Miscellaneous": [row("BFS01", "Gloves", 2.0)]})
    assert item("BFS99").is_active


# --- Sections and setup --------------------------------------------------------

def test_sections_from_heading_rows_and_flammable_flag():
    wb = Workbook()
    ws = wb.active
    ws.title = "All items"
    ws.append(["BFS Catalogue"])
    ws.append(HEADERS)
    ws.append(["Flammable Items"])
    ws.append(row("FLA001", "White spirit 1L", 3.0))
    ws.append(["Kitchens"])
    ws.append(row("BK001", "Basin 500mm", 40.0))
    ws.append(row("BK002", "Black granite bullnose worktop 3m", 90.0))
    buf = io.BytesIO()
    wb.save(buf)
    run = importer.preview_upload("all.xlsx", buf.getvalue(), None)
    importer.apply_run(run, None)
    assert item("FLA001").section_id == Section.FLAMMABLE and item("FLA001").is_flammable
    assert item("BK001").section_id == Section.KITCHENS and not item("BK001").is_flammable
    worktop = item("BK002")
    assert worktop.measure_type == MeasureType.CUT_TO_ORDER and worktop.catalogue_length_m == Decimal("3.00")


def test_recorded_item_setup_decisions():
    _, report = run_import({
        "Plumbing": [row("PLU037", "Sink clips", 2.0), row("PLU095", "Pipe clips", 5.0),
                     row("PLU020", "Solvent cement 250ml", 4.0)],
        "Plastering": [row("PLA007", "Angle bead 2.4m", 1.5), row("PLA015", "Scrim tape 90m", 2.0)],
        "Flooring": [row("V0002", "Cove former", 3.0), row("V0004", "Weld rod", 8.0),
                     row("V0020", "Safety vinyl", 18.0, unit="m2"), row("V0021", "Vinyl sheet 2m wide", 15.0)],
        "Misc": [row("BFS05", "Wipes", 6.0)],
    })
    assert (item("PLU037").measure_type, item("PLU037").pack_size) == (MeasureType.PACK, 8)
    assert (item("BFS05").measure_type, item("BFS05").pack_size) == (MeasureType.PACK, 80)
    assert item("PLA007").measure_type == MeasureType.WHOLE and item("PLA015").unit == "roll"
    assert item("V0002").measure_type == MeasureType.EACH and item("V0002").unit == "2m length"
    assert item("V0004").measure_type == MeasureType.EACH
    assert item("V0020").measure_type == MeasureType.AREA and item("V0020").unit == "m²"
    assert item("V0021").measure_type == MeasureType.AREA
    solvent = item("PLU020")
    assert solvent.section_id == Section.PLUMBING and not solvent.is_flammable
    notes = {x["part_no"]: x["note"] for x in report["setup_notes"]}
    assert "pack size" in notes["PLU095"] and "80" in notes["BFS05"]


def test_csv_in_windows_encoding():
    data = "Part No.,Description,Trade Price,Section\r\nADH014,Silicone 270ml white £,4.20,Adhesives\r\n".encode("cp1252")
    run = importer.preview_upload("adhesives.csv", data, None)
    importer.apply_run(run, None)
    assert item("ADH014").catalogue_name == "Silicone 270ml white £"


# --- Screens and permissions -------------------------------------------------

def test_import_screens_admin_only(client, make_user):
    client.force_login(make_user())
    assert client.get(reverse("catalogue:import_start")).status_code == 403


def test_upload_preview_apply_through_website(admin_client):
    upload = io.BytesIO(xlsx({"Adhesives": [row("ADH001", "Grab adhesive", 5.00), row("ADH001", "Dupe", 1.0)]}))
    upload.name = "catalogue.xlsx"
    resp = admin_client.post(reverse("catalogue:import_start"), {"file": upload})
    assert resp.status_code == 302
    page = admin_client.get(resp.url).content.decode()
    assert "Nothing has been changed yet" in page and "ADH001-B" in page
    assert Item.objects.count() == 0
    run = ImportRun.objects.get()
    admin_client.post(reverse("catalogue:import_apply", args=[run.pk]))
    assert Item.objects.count() == 2
    assert AuditEntry.objects.filter(action="catalogue.import").exists()


def test_bad_file_gives_clear_error(admin_client):
    upload = io.BytesIO(b"hello,world\n1,2\n")
    upload.name = "junk.csv"
    resp = admin_client.post(reverse("catalogue:import_start"), {"file": upload})
    assert resp.status_code == 200 and "No items found" in resp.content.decode()


def test_admin_edit_of_display_name_is_audited(admin_client, admin_user):
    run_import({"Adhesives": [row("ADH014", "Silicone sealant 270ml white C4 general purpose", 4.20)]})
    silicone = item("ADH014")
    url = reverse("admin:catalogue_item_change", args=[silicone.pk])
    form = admin_client.get(url).context["adminform"].form
    data = {k: v for k, v in form.initial.items() if v is not None and k in form.fields}
    data.update({"display_name": "Multi-purpose silicone 270ml white", "section": Section.ADHESIVES,
                 "is_active": "on", "pack_size": 1})
    data.pop("is_flammable", None)
    data.pop("code_to_confirm", None)
    data.pop("name_check", None)
    resp = admin_client.post(url, data)
    assert resp.status_code == 302, resp.context["adminform"].form.errors
    silicone.refresh_from_db()
    assert silicone.display_name == "Multi-purpose silicone 270ml white"
    assert silicone.catalogue_name == "Silicone sealant 270ml white C4 general purpose"
    entry = AuditEntry.objects.get(action="item.edit", target_id=str(silicone.pk))
    assert entry.actor == admin_user
    assert entry.changes["display name"]["to"] == "Multi-purpose silicone 270ml white"


def test_admin_editing_part_no_confirms_temporary_code(admin_client):
    run_import({"Electrical": [row("ELE118", "Twin socket", 4.10), row("ELE118", "Single socket", 3.20)]})
    temp = item("ELE118-B")
    url = reverse("admin:catalogue_item_change", args=[temp.pk])
    form = admin_client.get(url).context["adminform"].form
    data = {k: v for k, v in form.initial.items() if v is not None and k in form.fields}
    data.update({"part_no": "ELE130", "is_active": "on", "code_to_confirm": "on"})
    data.pop("is_flammable", None)
    data.pop("name_check", None)
    resp = admin_client.post(url, data)
    assert resp.status_code == 302
    temp.refresh_from_db()
    assert temp.part_no == "ELE130" and not temp.code_to_confirm


def test_cut_to_order_items_need_a_catalogue_length(admin_client):
    run_import({"Kitchens": [row("BK010", "Black granite bullnose worktop 3m", 90.0)]})
    worktop = item("BK010")
    url = reverse("admin:catalogue_item_change", args=[worktop.pk])
    form = admin_client.get(url).context["adminform"].form
    data = {k: v for k, v in form.initial.items() if v is not None and k in form.fields}
    data.update({"is_active": "on", "catalogue_length_m": ""})
    for key in ("is_flammable", "code_to_confirm", "name_check"):
        data.pop(key, None)
    resp = admin_client.post(url, data)
    assert resp.status_code == 200
    assert "catalogue_length_m" in resp.context["adminform"].form.errors


@pytest.mark.parametrize("filename,section", [
    ("f7411c88-Kitchens.csv", Section.KITCHENS),
    ("451e36e6-Flammable_Items.csv", Section.FLAMMABLE),
    ("ece420e9-Flooring.csv", Section.FLOORING),
    ("10d62a05-Miscellaneous.csv", Section.MISC),
])
def test_section_taken_from_uploaded_file_name(filename, section):
    data = "Part No.,Name,Trade Price\r\nX001,Thing,1.00\r\n".encode()
    importer.apply_run(importer.preview_upload(filename, data, None), None)
    assert item("X001").section_id == section


@pytest.mark.parametrize("code,name,section,expected", [
    ("ELE111", "UNIVOLT SAE16/25 TRUNKING EXTERNAL ANGLE 16 x 25", Section.ELECTRICAL, ("each", 1, "each")),
    ("ELE043", "APPLEBY SB619 1-GANG DRY LINING INSTALLATION BOX 35MM", Section.ELECTRICAL, ("each", 1, "each")),
    ("ELE058", "KRIMPTERM VIAS VCT8-W CABLE TIES 300 x 4.8mm (PACK 100) NATURAL", Section.ELECTRICAL, ("pack", 100, "pack")),
    ("ELE039", "UNICRIMP FIRE CLIP FOR MINI TRUNKING 25mm PK10", Section.ELECTRICAL, ("pack", 10, "pack")),
    ("ELE067", "UNIVOLT MIKA2 SELF ADHESIVE TRUNKING 16 x 25mm PER MTR", Section.ELECTRICAL, ("each", 1, "metre")),
    ("JOI009", "COARSE 55MM COLLATED DRYWALL SCREW PK 1000", Section.JOINERY, ("pack", 1000, "pack")),
    ("JOI086", "OWLETT FINE THREAD DRYWALL SCREW ZINC YELLOW 25MM (BOX 1000)", Section.JOINERY, ("pack", 1000, "pack")),
    ("JOI113", "SELF DRILL 38mm BULK DRYWALL SCREWPER 1000", Section.JOINERY, ("pack", 1000, "pack")),
    ("JOI103", "GALV COLLATED BRADS F16X32MM (BOX2000)", Section.JOINERY, ("pack", 2000, "pack")),
    ("JOI059", "METAL ECLIPSE 34MM CUTTING MULTI TOOL BLADE PK OF 5", Section.JOINERY, ("pack", 5, "pack")),
    ("JOI101", "Blue Spot Tools 29190 Utility Blade Set in Dispenser (50 Pieces)", Section.JOINERY, ("pack", 50, "pack")),
    ("JOI080", "Black Granite Effect Bullnose Worktop 3m 38mm/8mm (600 Deep)", Section.JOINERY, ("cut", 1, "m")),
    ("JOI081", "Black Worktop Joint Strip 38mm/6mm/8mm", Section.JOINERY, ("each", 1, "each")),
    ("JOI083", "Black Worktop End Cap 38mm/6mm/8mm", Section.JOINERY, ("each", 1, "each")),
    ("PLA003", "BRITISH GYPSUM THISTLE BONDING PLASTER BAG 25kg (CARLITE) 06055/7", Section.PLASTERING, ("each", 1, "each")),
    ("V0003", "POLYSAFE HYDRO 4930 WHITE STONE", Section.FLOORING, ("area", 1, "m²")),
    ("V0002", "Cove Former BLACK 2M LENGTH X 40 BOX", Section.FLOORING, ("each", 1, "2m length")),
])
def test_setup_rules_on_real_catalogue_names(code, name, section, expected):
    from catalogue.setup_rules import initial_setup

    fields, _ = initial_setup(code, name, "", section)
    assert (fields["measure_type"], fields["pack_size"], fields["unit"]) == expected
